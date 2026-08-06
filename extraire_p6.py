"""Extrait du planning P6 les activités qui alimentent le Gantt, et elles seules.

    python3 extraire_p6.py TUXERev2LinkedActivities.xlsx donnees_p6.txt

Le solveur n'en voit rien : ces dates ne servent qu'à remplir des lignes du planning
d'occupation que la simulation ne calcule pas — travaux marins et finitions. Le plan de
production, lui, reste entièrement issu du solveur.

Une catégorie donne une barre par élément : la plus précoce des dates de début de ses
activités, la plus tardive de leurs dates de fin. C'est le niveau de WBS demandé, pas plus
fin : le Gantt porte l'occupation d'une ressource, pas le détail des tâches.
"""
import re, sys
from collections import defaultdict
from openpyxl import load_workbook

SRC, DST = sys.argv[1], sys.argv[2]

# Catégories retenues. Les finitions se repèrent par leur nom de WBS — c'est la maille
# demandée ; les travaux marins par le nom d'activité, faute d'un WBS qui les isole.
PAR_WBS = {
    'Omega seal installation': 'omega',
    'Immersion Joint Removal': 'bhrem',          # dépose des bulkheads au joint d'immersion
    'Immersion Joint infill Concrete': 'infill',
    'Removal of TE System': 'tesys',
    'Drainage Installation': 'drain',
    'Walkways': 'walk',
}
PAR_NOM = [
    (re.compile(r'^Gravel Bed', re.I), 'gravel'),
    (re.compile(r'Leveling Layer Installation', re.I), 'level'),
    (re.compile(r'^(Locking fill|Backfill)', re.I), 'lock'),
    (re.compile(r'Element ready for floatup', re.I), 'readyfu'),
]

# Identifiant d'élément dans un nom d'activité. Le planning en use plusieurs formes —
# « TE-59 », « TE59 », « SP-01 », « SP10 », « SPE.10 » — et les travaux de joint portent deux
# numéros (« TE-77-78 ») : c'est le premier qui désigne l'élément concerné.
RX_ID = re.compile(r'\b(S?TE|SPE?)[\s.\-]*0*(\d{1,2})\b', re.I)

def identifiant(nom):
    m = RX_ID.search(nom)
    if not m:
        return None
    fam = m.group(1).upper()
    n = int(m.group(2))
    if fam.startswith('SP'):
        return 'SPE-%02d' % n
    return 'STE-%02d' % n

def jour(v):
    return v.date().isoformat() if hasattr(v, 'date') else None

ws = load_workbook(SRC, read_only=True, data_only=True)['TASK']
bornes = defaultdict(lambda: [None, None])
lues = ignorees = 0
for r in ws.iter_rows(min_row=3, values_only=True):
    nom = r[3]
    if nom is None:
        continue
    nom = str(nom)
    wbs = str(r[9] or '')
    cat = PAR_WBS.get(wbs)
    if not cat:
        for rx, c in PAR_NOM:
            if rx.search(nom):
                cat = c
                break
    if not cat:
        continue
    ident = identifiant(nom)
    d, f = jour(r[4]), jour(r[5])
    if not ident or not d or not f:
        ignorees += 1
        continue
    lues += 1
    b = bornes[(ident, cat)]
    if b[0] is None or d < b[0]:
        b[0] = d
    if b[1] is None or f > b[1]:
        b[1] = f

lignes = ['%s|%s|%s|%s' % (i, c, b[0], b[1]) for (i, c), b in sorted(bornes.items())]
open(DST, 'w').write(';'.join(lignes))

par_cat = defaultdict(int)
for (i, c) in bornes:
    par_cat[c] += 1
print(f'{DST} écrit — {lues} activités retenues, {ignorees} sans élément ou sans date, '
      f'{len(bornes)} barres')
for c in sorted(par_cat):
    print(f'   {c:8} : {par_cat[c]:3} éléments')
