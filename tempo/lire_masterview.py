"""Lit `MASTERVIEW.xlsm`, le classeur de consolidation logistique (colisage, poids,
camions, fenêtres de livraison) — la donnée qui manquait pour tout le volet logistique
signalé dans `tempo/DONNEES.md` (le « Jules Tab » évoqué en réunion, en plus détaillé).

    python3 tempo/lire_masterview.py MASTERVIEW.xlsm

Le classeur mélange des tables de travail (sous-ensembles d'éléments, brouillons) et une
table consolidée. Ce module ne lit que ce qui est identifié comme la version consolidée —
voir chaque fonction pour la feuille exacte et pourquoi elle a été retenue plutôt qu'une
autre qui lui ressemble.
"""
import argparse
import warnings

from openpyxl import load_workbook

warnings.filterwarnings('ignore', module='openpyxl')

COLONNES_COLIS = ['code', 'element', 'plan', 'poids_final', 'fournisseur', 'designation',
                   'poids_kg', 'tun', 'date_livraison', 'type', 'categorie', 'source_commande',
                   'remarques', 'couleur', 'controle_qualite']


def lire_colis(chemin, feuille='ALL ELEMENTS'):
    """Le registre des colis (un colis = une pièce livrée pour un élément donné) : poids,
    fournisseur, date de livraison, catégorie (BS/WA/TS/2nde phase).

    `ALL ELEMENTS` est retenue plutôt que `Data`, `Sheet1`, `Sheet1 (2)`, `Sheet1 (3)` —
    quatre autres feuilles qui portent la même structure de colonnes mais chacune sur un
    sous-ensemble d'éléments (36, 27, 2 et 19 éléments distincts respectivement, contre
    84 sur `ALL ELEMENTS` — la somme des quatre, à l'élément près) : ce sont des feuilles
    de travail, `ALL ELEMENTS` en est la consolidation.
    """
    wb = load_workbook(chemin, data_only=True)
    ws = wb[feuille]
    out = []
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 1).value
        if not code:
            continue
        vals = [ws.cell(r, c).value for c in range(1, len(COLONNES_COLIS) + 1)]
        out.append(dict(zip(COLONNES_COLIS, vals)))
    return out


COLONNES_LIVRAISON = ['lignes', 'element_segment', 'delivery_id', 'plateforme', 'segment',
                        'partie', 'designation', 'evenement', 'date_debut', 'au_plus_tot',
                        'au_plus_tard', 'date_retenue', 'lieu_livraison', 'temps_chargement',
                        'coefficient_semaine']


def lire_livraisons(chemin, feuille='DeliveryPlan'):
    """Le plan de livraison : une ligne par livraison, avec sa fenêtre (au plus tôt / au
    plus tard) et la date effectivement retenue — la donnée qui permettrait de calibrer
    le calendrier absolu du moteur (`tempo/moteur/`) sur de vraies dates plutôt que sur
    les deux hypothèses posées dans `tempo/moteur/ARCHITECTURE.md`."""
    wb = load_workbook(chemin, data_only=True)
    ws = wb[feuille]
    out = []
    for r in range(2, ws.max_row + 1):
        if not ws.cell(r, 1).value:
            continue
        vals = [ws.cell(r, c).value for c in range(1, len(COLONNES_LIVRAISON) + 1)]
        out.append(dict(zip(COLONNES_LIVRAISON, vals)))
    return out


COLONNES_CAMION = ['plateforme', 'segment', 'partie', 'set_', 'sous_set', 'nb_camions',
                    'avec_rack_blanc', 'designation', 'mode', 'evenement', 'poids_kg',
                    'lieu_livraison', 'au_plus_tot', 'au_plus_tard']


def lire_camions(chemin, feuille='LIST'):
    """Le détail camion par camion : nombre de camions par sous-lot, présence d'un
    « White Rack » (le rack « mesh », cf. réunion du 10/08 — Ø marque son absence, donc
    probablement une livraison en vrac plutôt qu'en rack), mode de livraison (LOOSE /
    ASSEMBLY / BOTH), poids, fenêtre de livraison en jours relatifs à la coulée."""
    wb = load_workbook(chemin, data_only=True)
    ws = wb[feuille]
    out = []
    for r in range(2, ws.max_row + 1):
        if not ws.cell(r, 1).value:
            continue
        vals = [ws.cell(r, c).value for c in range(1, len(COLONNES_CAMION) + 1)]
        out.append(dict(zip(COLONNES_CAMION, vals)))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('classeur')
    args = ap.parse_args()

    colis = lire_colis(args.classeur)
    dates = [c['date_livraison'] for c in colis if hasattr(c['date_livraison'], 'year')]
    print(f"{len(colis)} colis (feuille ALL ELEMENTS), "
          f"{len({c['element'] for c in colis})} éléments distincts.")
    if dates:
        print(f"  dates de livraison : {min(dates)} -> {max(dates)}")
    categories = {}
    for c in colis:
        categories[c['categorie']] = categories.get(c['categorie'], 0) + 1
    print(f"  répartition par catégorie : {categories}")

    livraisons = lire_livraisons(args.classeur)
    print(f"\n{len(livraisons)} livraisons planifiées (feuille DeliveryPlan).")
    lignes = {}
    for l in livraisons:
        lignes[l['lignes']] = lignes.get(l['lignes'], 0) + 1
    print(f"  par ligne : {lignes}")

    def nombre(v):
        return v if isinstance(v, (int, float)) else 0

    camions = lire_camions(args.classeur)
    print(f"\n{len(camions)} lots camion (feuille LIST).")
    total_camions = sum(nombre(c['nb_camions']) for c in camions)
    print(f"  {total_camions} camions au total, poids total "
          f"{sum(nombre(c['poids_kg']) for c in camions) / 1000:.0f} t.")


if __name__ == '__main__':
    main()
