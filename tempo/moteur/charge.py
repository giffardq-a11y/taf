"""Construit la charge de ressource agrégée, jour calendaire absolu par jour calendaire
absolu, en recopiant le gabarit de tâches N3 sur les 5 lignes déphasées du N1 — voir
`tempo/moteur/ARCHITECTURE.md` pour le principe et les deux hypothèses posées.

    python3 -m tempo.moteur.charge General_Tempo_Staggering.xlsx STE__TEMPO__N3_Takt_Plan__V0_1.xlsx
"""
import argparse
import pathlib
import re
import sys
from collections import defaultdict

# Permet `python3 tempo/moteur/charge.py` directement, sans installation du paquet :
# la racine du dépôt (deux niveaux au-dessus de ce fichier) doit être sur sys.path.
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2]))
from tempo.lire_tempo import lire_n1, lire_taches_n3
from tempo.moteur.modele import ParametresCalage, PointDeCharge


def params_depuis_n1(n1):
    """Construit `ParametresCalage.demarrage_tempo` à partir de la colonne `demarrage`
    de `lire_n1` (ex. « L4 - T25 » -> {'Line 4 - Option 343': 25}). Une ligne peut
    apparaître plusieurs fois dans `n1` (une par catégorie) ; on ne garde le
    déphasage qu'une fois par ligne."""
    demarrage = {}
    for e in n1:
        if e['ligne'] in demarrage:
            continue
        m = re.search(r'T(\d+)', str(e['demarrage'] or ''))
        if m:
            demarrage[e['ligne']] = int(m.group(1))
    return ParametresCalage(demarrage_tempo=demarrage)


def construire_charge(taches, charges, params, horizon_jours=140):
    """Renvoie {(ressource, jour_absolu, fenetre): PointDeCharge}.

    Rejoue le gabarit N3 (commun aux 5 lignes) sur chaque ligne, décalé de son
    déphasage (`params.demarrage_tempo`), répété tous les `periode_relance_jours`
    autant de fois qu'il en tient dans `horizon_jours` — assez pour voir le motif de
    charge se répéter en régime établi, pas seulement au premier élément.
    """
    par_cle = {(t['feuille'], t['ligne_excel']): t for t in taches}
    points = {}

    for ligne, demarrage in params.demarrage_tempo.items():
        nb_elements = max(1, horizon_jours // params.periode_relance_jours)
        for indice in range(nb_elements):
            offset = demarrage + indice * params.periode_relance_jours - 1  # Ti=1 -> jour démarrage
            for c in charges:
                if c['jalon'] or not c['effectif'] or c['tempo'] is None:
                    continue
                t = par_cle.get((c['feuille'], c['ligne_excel']))
                ressource = (t and t.get('ressource')) or 'NC (type non renseigné)'
                jour_absolu = offset + c['tempo']
                cle = (ressource, jour_absolu, c['fenetre'])
                if cle not in points:
                    points[cle] = PointDeCharge(ressource=ressource, jour_absolu=jour_absolu,
                                                 fenetre=c['fenetre'], demande=0)
                p = points[cle]
                p.demande += c['effectif']
                nom_tache = (t and t.get('nom')) or '?'
                p.detail.append((ligne, f"{c['zone']}/{nom_tache}", c['effectif']))
    return points


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('n1')
    ap.add_argument('n3')
    ap.add_argument('--horizon', type=int, default=140, help='jours simulés (défaut 140)')
    ap.add_argument('--ressource', default=None, help='ne montrer que cette ressource')
    ap.add_argument('--top', type=int, default=15, help='nombre de pics à afficher')
    args = ap.parse_args()

    from openpyxl import load_workbook
    import warnings
    warnings.filterwarnings('ignore', module='openpyxl')
    wb1 = load_workbook(args.n1, data_only=True)
    feuille_n1 = sorted([n for n in wb1.sheetnames if '5 Lines' in n])[-1]
    n1 = lire_n1(args.n1, feuille_n1)
    params = params_depuis_n1(n1)
    print(f"Déphasage retenu (feuille '{feuille_n1}') : {params.demarrage_tempo}")

    taches, charges = lire_taches_n3(args.n3)
    points = construire_charge(taches, charges, params, args.horizon)

    par_ressource = defaultdict(list)
    for p in points.values():
        if args.ressource and p.ressource != args.ressource:
            continue
        par_ressource[p.ressource].append(p)

    print(f"\n{len(par_ressource)} ressource(s) sollicitée(s) sur un horizon de "
          f"{args.horizon} jours, {len(points)} points de charge (ressource × jour × fenêtre).")
    for ressource, pts in sorted(par_ressource.items(), key=lambda kv: -max(p.demande for p in kv[1])):
        pts.sort(key=lambda p: -p.demande)
        pic = pts[0]
        print(f"\n=== {ressource} — pic {pic.demande:.1f} au jour {pic.jour_absolu}, "
              f"fenêtre {pic.fenetre} ===")
        for p in pts[:args.top]:
            lignes = sorted({d[0] for d in p.detail})
            print(f"  jour {p.jour_absolu:>4}  fenêtre {str(p.fenetre):>14}  "
                  f"demande={p.demande:>6.1f}  lignes concernées: {', '.join(lignes)}")


if __name__ == '__main__':
    main()
