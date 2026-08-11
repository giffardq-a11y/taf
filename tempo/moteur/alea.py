"""Premier module de réponse à un aléa — retard ou panne matérielle sur une ressource.

Ce n'est **pas** un réordonnancement complet : ça suppose de connaître l'ordonnancement
détaillé tâche par tâche (les liens de précédence, pas encore construits — voir
`tempo/moteur/ARCHITECTURE.md`, section « ce qui n'est pas encore construit »). Ce que ce
module fait, dans la limite de ce que la charge déjà calculée permet : évaluer l'impact
d'une perte de capacité sur une ressource, pendant une fenêtre donnée — le déficit que ça
crée, et le tampon minimal qui suffirait à l'absorber sans aucun jour en déficit. Répond
à l'objectif « replanifier les ressources en cas de retard de casse matériel » posé au
départ du simulateur, au niveau où le moteur peut y répondre aujourd'hui.

    python3 -m tempo.moteur.alea <n1.xlsx> <n3.xlsx> --ressource BC --zone S \\
        --capacite 85 --perte 20 --jour-debut 100 --duree 3
"""
import argparse
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[2]))


def simuler_perte_capacite(points_ressource, capacite_nominale, perte, jour_debut, jour_fin):
    """`points_ressource` : liste de `PointDeCharge` d'UNE SEULE ressource (filtrer en
    amont). `jour_debut`/`jour_fin` doivent être du même type que `PointDeCharge.jour_absolu`
    (un entier pour la main-d'œuvre via `charge.py`, un `datetime.date` pour la logistique
    via `logistique.py`) — le module ne suppose rien de plus qu'un ordre total sur ce type.

    Renvoie un rapport : le déficit qui existait déjà avant l'aléa (`deficit_prealable`,
    demande > capacité nominale — un aléa ne crée pas un problème qui existait déjà, il
    l'aggrave), le déficit pendant l'incident (`deficit_incident`, à capacité réduite), et
    le tampon minimal qui ramènerait le déficit incident à zéro (`tampon_minimal` — la
    plus grande insuffisance rencontrée pendant la fenêtre, jamais moins que ce qu'il faut
    pour le pire jour)."""
    capacite_effective = max(0.0, capacite_nominale - perte)
    touches = [p for p in points_ressource if jour_debut <= p.jour_absolu < jour_fin]
    touches.sort(key=lambda p: (p.jour_absolu, str(p.fenetre)))

    detail = []
    deficit_prealable_max = 0.0
    deficit_incident_max = 0.0
    for p in touches:
        deficit_prealable = max(0.0, p.demande - capacite_nominale)
        deficit_incident = max(0.0, p.demande - capacite_effective)
        deficit_prealable_max = max(deficit_prealable_max, deficit_prealable)
        deficit_incident_max = max(deficit_incident_max, deficit_incident)
        detail.append({
            'jour': p.jour_absolu, 'fenetre': p.fenetre, 'demande': p.demande,
            'deficit_prealable': deficit_prealable, 'deficit_incident': deficit_incident,
        })

    return {
        'capacite_nominale': capacite_nominale,
        'perte': perte,
        'capacite_effective': capacite_effective,
        'nb_points_touches': len(touches),
        'deficit_prealable_max': deficit_prealable_max,
        'deficit_incident_max': deficit_incident_max,
        'tampon_minimal': deficit_incident_max,  # unités de capacité à restaurer pour repasser à 0 déficit
        'detail': detail,
    }


def rapport_texte(rapport, ressource):
    lignes = []
    lignes.append(f"Ressource {ressource} — capacité nominale {rapport['capacite_nominale']:.0f}, "
                   f"perte simulée {rapport['perte']:.0f} -> capacité effective "
                   f"{rapport['capacite_effective']:.0f}")
    lignes.append(f"{rapport['nb_points_touches']} points de charge dans la fenêtre de l'incident.")
    if rapport['deficit_prealable_max'] > 0:
        lignes.append(f"⚠ Déficit déjà présent SANS l'aléa (demande > capacité nominale) : "
                       f"jusqu'à {rapport['deficit_prealable_max']:.1f} — l'aléa aggrave un problème "
                       f"préexistant, il ne le crée pas seul.")
    else:
        lignes.append("Aucun déficit sans l'aléa : la capacité nominale suffisait sur cette fenêtre.")
    lignes.append(f"Déficit pendant l'incident (capacité réduite) : jusqu'à "
                   f"{rapport['deficit_incident_max']:.1f}.")
    lignes.append(f"Tampon minimal pour ramener le déficit à zéro pendant l'incident : "
                   f"{rapport['tampon_minimal']:.1f} unités de {ressource}.")
    return "\n".join(lignes)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                  formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('n1')
    ap.add_argument('n3')
    ap.add_argument('--ressource', required=True, help='code ressource, ex. BC')
    ap.add_argument('--zone', default=None, help="ne garder que les charges de cette zone N3 (ex. S)")
    ap.add_argument('--capacite', type=float, required=True, help='capacité nominale de la ressource')
    ap.add_argument('--perte', type=float, required=True, help='unités perdues pendant l\'incident')
    ap.add_argument('--jour-debut', type=int, required=True, help='jour_absolu de début de l\'incident')
    ap.add_argument('--duree', type=int, default=1, help='durée en jours (défaut 1)')
    ap.add_argument('--horizon', type=int, default=140)
    args = ap.parse_args()

    import warnings
    warnings.filterwarnings('ignore', module='openpyxl')
    from openpyxl import load_workbook
    from tempo.lire_tempo import lire_n1, lire_taches_n3
    from tempo.moteur.charge import construire_charge, params_depuis_n1

    wb1 = load_workbook(args.n1, data_only=True)
    feuille_n1 = sorted([n for n in wb1.sheetnames if '5 Lines' in n])[-1]
    n1 = lire_n1(args.n1, feuille_n1)
    params = params_depuis_n1(n1)
    taches, charges = lire_taches_n3(args.n3)
    if args.zone:
        charges = [c for c in charges if c['zone'] == args.zone]
    points = construire_charge(taches, charges, params, args.horizon)

    points_ressource = [p for p in points.values() if p.ressource == args.ressource]
    if not points_ressource:
        print(f"Aucun point de charge pour la ressource {args.ressource!r} "
              f"{'en zone ' + args.zone if args.zone else ''}.")
        return

    rapport = simuler_perte_capacite(points_ressource, args.capacite, args.perte,
                                       args.jour_debut, args.jour_debut + args.duree)
    print(rapport_texte(rapport, args.ressource))
    if rapport['detail']:
        print("\nDétail par point de charge touché :")
        for d in rapport['detail']:
            print(f"  jour {d['jour']:>4}  fenêtre {str(d['fenetre']):>14}  demande={d['demande']:>6.1f}  "
                  f"déficit sans aléa={d['deficit_prealable']:>6.1f}  déficit pendant l'incident="
                  f"{d['deficit_incident']:>6.1f}")


if __name__ == '__main__':
    main()
