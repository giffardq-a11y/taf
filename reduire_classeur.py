"""Produit un classeur réduit : uniquement les feuilles, colonnes et lignes que le solveur lit.

Ce que le solveur lit, relevé dans parseWorkbook (index.html) :
  Inputs         : col A = n° d'ordre ; colonnes PL-1..SPE = identifiants, chacune suivie
                   de sa colonne de statut as-built.
  Immersion      : ID Element, Date Immersion, Activity Name, Start, Finish. Rien d'autre.
  Config_Cycles  : colonnes A/B/C (ligne, date seuil 1, cycle 1). Les seuils 2 à 4 sont
                   déterminés par le solveur, jamais lus.
  Config_SPE     : colonnes A à G (zone, nom, activité, durée, type, jalon étanchéité,
                   rythme SPE long).
Feuilles ignorées par le solveur : Feuil1, Planning_Final, Dashboard, Recap_Dates,
Config_Outfitting (lue mais jamais appliquée).
"""
import re, sys
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

SRC, DST = sys.argv[1], sys.argv[2]
src = load_workbook(SRC, data_only=True)
out = Workbook()
out.remove(out.active)
POLICE = Font(name='Arial', size=10)
GRAS = Font(name='Arial', size=10, bold=True)

def poser(ws, lignes, entetes=0):
    for r, ligne in enumerate(lignes, start=1):
        for c, v in enumerate(ligne, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = GRAS if r <= entetes else POLICE
            if hasattr(v, 'year'):
                cell.number_format = 'DD/MM/YYYY'

# ---- Inputs : n° d'ordre + (identifiant, statut) par ligne de production ----
si = src['Inputs']
brut = [[c.value for c in row] for row in si.iter_rows()]
hdr = next(i for i, r in enumerate(brut) if any(str(v).strip() == 'PL-1' for v in r if v is not None))
colonnes = {}
for nom in ('PL-1', 'PL-2', 'PL-3', 'PL-4', 'PL-5', 'SPE'):
    colonnes[nom] = next(i for i, v in enumerate(brut[hdr]) if str(v).strip() == nom)

lignes = [['Sequence Production'], []]
entete = ['No.']
for nom in colonnes:
    entete += [nom, 'Statut']
lignes.append(entete)
gardees = 0
for r in brut[hdr + 1:]:
    no = r[0] if r else None
    if not isinstance(no, (int, float)):
        continue
    sortie = [int(no)]
    for nom, c in colonnes.items():
        ident = str(r[c] or '').strip()
        statut = str(r[c + 1] or '').strip() if c + 1 < len(r) else ''
        sortie += [ident or None, statut or None]
    lignes.append(sortie)
    gardees += 1
poser(out.create_sheet('Inputs'), lignes, entetes=3)

# ---- Immersion : 5 colonnes, et seules les lignes qui portent une information ----
sm = src['Immersion']
tete = [str(c.value).strip() if c.value is not None else '' for c in sm[1]]
idx = {nom: tete.index(nom) for nom in ('ID Element', 'Date Immersion', 'Activity Name', 'Start', 'Finish')}
# Les quatre activités que le solveur cherche par leur nom, avec les motifs exacts qu'il
# emploie : une activité qui n'y répond pas n'entre dans aucun calcul.
MOTIFS = (re.compile(r'Immersion\s*-\s*(TE|SP)-?\s?(\d+)', re.I),
          re.compile(r'clamp.*SP-?\s?(\d+)|SP-?\s?(\d+).*clamp', re.I),
          re.compile(r'floatout.*Concrete\s*-\s*(TE|SPE?)-?\s?(\d+)', re.I),
          re.compile(r'ready.*float|float.*up.*ready', re.I))

immersion = [['ID Element', 'Date Immersion', 'Activity Name', 'Start', 'Finish']]
lues = 0
for row in sm.iter_rows(min_row=2):
    vals = [c.value for c in row]
    def val(nom):
        i = idx[nom]
        return vals[i] if i < len(vals) else None
    lues += 1
    ident = str(val('ID Element') or '').strip()
    act = val('Activity Name')
    acte = act if isinstance(act, str) else ''
    # Niveau hiérarchique du planning P6 : l'indentation porte la structure. Niveau <= 4 =
    # la zone, niveau 6 = le groupe du SPE dans cette zone. Au-delà, seules comptent les
    # activités que le solveur cherche par leur nom.
    niveau = len(acte) - len(acte.lstrip())
    utile = bool(ident) or (acte.strip() and niveau <= 6) or any(m.search(acte) for m in MOTIFS)
    if not utile:
        continue
    immersion.append([val('ID Element'), val('Date Immersion'), act, val('Start'), val('Finish')])
poser(out.create_sheet('Immersion'), immersion, entetes=1)

# ---- Config_Cycles : seuls Date Seuil 1 / Cycle 1 sont lus ----
sc = src['Config_Cycles']
cycles = [['Ligne (PL)', 'Date Seuil 1', 'Cycle 1 (sem)']]
for r in range(2, 7):
    cycles.append([sc.cell(r, 1).value, sc.cell(r, 2).value, sc.cell(r, 3).value])
poser(out.create_sheet('Config_Cycles'), cycles, entetes=1)

# ---- Config_SPE : les 7 zones de la chaîne SPE ----
ss = src['Config_SPE']
spe = [['Ndeg Zone', 'Nom Zone', 'Activite / Phase', 'Duree fixe (semaines)',
        'Type (Beton ou Outfitting)', 'jalon etancheite (sem)', 'rythme SPE long (sem)']]
for r in range(2, 9):
    spe.append([ss.cell(r, c).value for c in range(1, 8)])
poser(out.create_sheet('Config_SPE'), spe, entetes=1)

for ws in out.worksheets:
    for c in range(1, ws.max_column + 1):
        largeur = max((len(str(ws.cell(r, c).value or '')) for r in range(1, ws.max_row + 1)), default=8)
        ws.column_dimensions[get_column_letter(c)].width = min(max(largeur + 2, 9), 46)
    ws.freeze_panes = 'A2'

out.save(DST)
print(f"Inputs      : {gardees} lignes conservées, {1 + 2 * len(colonnes)} colonnes")
print(f"Immersion   : {len(immersion) - 1} lignes conservées sur {lues}, 5 colonnes sur {len(tete)}")
print(f"Config_Cycles : 5 lignes, 3 colonnes sur 9")
print(f"Config_SPE  : 7 zones, 7 colonnes")
